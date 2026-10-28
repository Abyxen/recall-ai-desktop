from pathlib import Path
import pytest
from recall.documents import extract


def test_docx(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph('Mitosis is cell division.')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Telophase'
    path = tmp_path / 'test.docx'
    doc.save(path)
    pages = extract(path)
    assert 'Mitosis' in pages[0].text and 'Telophase' in pages[0].text


def test_pptx(tmp_path):
    from pptx import Presentation
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = 'First slide'
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = 'Second slide'
    path = tmp_path / 'test.pptx'
    presentation.save(path)
    pages = extract(path)
    assert pages[1].number == 2 and 'Second slide' in pages[1].text


def test_pdf_and_encrypted(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(600, 800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (PDF source evidence) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    path = tmp_path / 'test.pdf'
    writer.write(path)
    assert 'PDF source evidence' in extract(path)[0].text
    writer.encrypt('password')
    writer.write(path)
    with pytest.raises(ValueError, match='Encrypted'):
        extract(path)


def test_markdown_and_utf16(tmp_path):
    path = tmp_path / 'notes.md'
    path.write_text('# Lesson\nHello\fSecond section', encoding='utf-16')
    pages = extract(path)
    assert pages[1].text == 'Second section'


def test_images_without_ocr(tmp_path):
    from PIL import Image
    path = tmp_path / 'image.png'
    Image.new('RGB', (20, 20)).save(path)
    with pytest.raises(ValueError, match='OCR'):
        extract(path)


def test_no_text_and_unsupported(tmp_path):
    path = tmp_path / 'blank.txt'
    path.write_text('   ')
    with pytest.raises(ValueError, match='No readable'):
        extract(path)
    path = tmp_path / 'bad.exe'
    path.write_text('test')
    with pytest.raises(ValueError, match='Unsupported'):
        extract(path)
