from __future__ import annotations

import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path

MAX_BYTES = 50 * 1024 * 1024
MAX_TEXT = 2_000_000
SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}


@dataclass
class Page:
    number: int
    text: str


def ocr_available():
    return bool(shutil.which("tesseract"))


def ocr_image(image):
    if not ocr_available():
        raise ValueError("OCR needs the optional Tesseract program on PATH. Install it, restart Recall, then retry.")
    import pytesseract
    return pytesseract.image_to_string(image, timeout=45)


def extract(path: Path, ocr=False) -> list[Page]:
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {path.suffix or '(none)'}")
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError("The file is missing or empty.")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Files must be 50 MB or smaller.")
    suffix = path.suffix.lower()
    try:
        if suffix in {".docx", ".pptx"}:
            with zipfile.ZipFile(path) as archive:
                if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                    raise ValueError("The Office archive expands beyond the 200 MB safety limit.")
        if suffix in {".txt", ".md"}:
            try:
                text = path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError:
                text = path.read_text(encoding="utf-16")
            pages = [Page(i + 1, text) for i, text in enumerate(text.split('\f'))]
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported. Import an unlocked copy.")
            if len(reader.pages) > 1000:
                raise ValueError("PDFs must contain at most 1,000 pages.")
            pages = [Page(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
            if ocr and any(not p.text.strip() for p in pages):
                try:
                    import pymupdf
                except ImportError:
                    raise ValueError("Scanned PDF OCR also needs optional PyMuPDF: pip install PyMuPDF. Then re-index with OCR enabled.")
                from PIL import Image
                with pymupdf.open(path) as pdf:
                    for page in pages:
                        if not page.text.strip():
                            pix = pdf[page.number - 1].get_pixmap(matrix=pymupdf.Matrix(2, 2))
                            page.text = ocr_image(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))
        elif suffix == ".docx":
            from docx import Document
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
            pages = [Page(1, text)]  # DOCX does not store reliable rendered page positions.
        elif suffix == ".pptx":
            from pptx import Presentation
            slides = Presentation(path).slides
            pages = []
            for i, slide in enumerate(slides):
                texts = [s.text for s in slide.shapes if s.has_text_frame]
                texts += [' | '.join(c.text for c in row.cells) for s in slide.shapes if s.has_table for row in s.table.rows]
                pages.append(Page(i + 1, '\n'.join(texts)))
        else:
            if not ocr:
                raise ValueError("Image import requires OCR. Enable OCR in Settings and install Tesseract.")
            from PIL import Image
            with Image.open(path) as image:
                if image.width * image.height > 40_000_000:
                    raise ValueError("Image exceeds the 40 megapixel limit.")
                pages = [Page(1, ocr_image(image))]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Cannot read {path.name}. It may be corrupt or in an unsupported encoding. {exc}") from exc
    if sum(len(p.text) for p in pages) > MAX_TEXT:
        raise ValueError("Extracted text exceeds the two million character limit. Split the document.")
    if not any(p.text.strip() for p in pages):
        raise ValueError("No readable text found. For scanned material, enable optional OCR in Settings.")
    return pages


def chunk_pages(pages: list[Page], size=180, overlap=35):
    if size <= 0 or not 0 <= overlap < size:
        raise ValueError("Chunk size must be positive and overlap smaller than size.")
    chunks = []
    for page in pages:
        words = re.sub(r"\s+", " ", page.text).strip().split()
        for start in range(0, len(words), size - overlap):
            chunks.append((page.number, ' '.join(words[start:start + size])))
            if start + size >= len(words):
                break
    return chunks
