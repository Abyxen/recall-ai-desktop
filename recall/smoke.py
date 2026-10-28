"""Small offline packaged-dependency check, using only isolated synthetic data."""
import tempfile
from pathlib import Path

from .db import Database
from .documents import extract
from .library import Library
from .retrieval import Retriever
from .study import Generator, review_card, save_cards, save_quiz, submit_quiz


def offline_workflow():
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        db = Database(root / 'data')
        text = 'Chlorophyll absorbs sunlight in chloroplasts during photosynthesis.'
        path = root / 'Synthetic.txt'
        path.write_text(text, encoding='utf-8')
        doc_id = Library(db).import_file(path)
        sources = Retriever(db).search('Chlorophyll sunlight')
        assert sources and sources[0]['document_id'] == doc_id
        raw = dict(question='What does chlorophyll absorb?', answer='sunlight', topic='Photosynthesis', difficulty='easy',
                   source_id=sources[0]['id'], quote=text, type='mcq', options=['sunlight', 'soil', 'DNA', 'glucose'], key_points=[], explanation=text)
        item = Generator.validate(raw, 'quiz', {sources[0]['id']: sources[0]}, ['mcq'])
        assert save_cards(db, [item], 'Smoke') == 1
        review_card(db, db.rows('SELECT id FROM cards')[0]['id'], 4)
        quiz_id = save_quiz(db, 'Smoke', [item])
        assert submit_quiz(db, quiz_id, ['sunlight'])['score'] == 1
        reopened = Database(root / 'data')
        assert len(reopened.rows('SELECT * FROM attempts')) == 1
        doc = Document()
        doc.add_paragraph(text)
        doc.save(root / 'Synthetic.docx')
        assert text in extract(root / 'Synthetic.docx')[0].text
        slides = Presentation()
        slides.slides.add_slide(slides.slide_layouts[1]).shapes.title.text = text
        slides.save(root / 'Synthetic.pptx')
        assert text in extract(root / 'Synthetic.pptx')[0].text
        writer = PdfWriter()
        page = writer.add_blank_page(600, 800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 50 700 Td (Synthetic PDF evidence) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(root / 'Synthetic.pdf')
        assert 'Synthetic PDF evidence' in extract(root / 'Synthetic.pdf')[0].text
    return True
